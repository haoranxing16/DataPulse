from openpyxl import load_workbook
from pathlib import Path
import csv
import json
import os
import tempfile
FileExtensionSamples = [".csv",".json",".xlsx"]

def csv_parse(filepath:Path)->list[dict]:
    records = []
    with filepath.open(newline="",encoding="utf-8") as f:
        for row in csv.DictReader(f):
            records.append(row)
    return records

def json_parse(filepath:Path)->list[dict]:
    with filepath.open(newline="",encoding="utf-8") as f:
        return json.load(f)

def xlsx_parse(filepath:Path)->list[dict]:
    wb = load_workbook(filepath,read_only=True,data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only = True)
    headers = next(rows)
    return [dict(zip(headers,row)) for row in rows]

def exceptdetection(filepath:Path):
    try:
        if not(filepath.exists() & filepath.is_file()):
            raise FileNotFoundError("The file is not exists or this path is not a file path")
    except FileNotFoundError as e:
        print(e)
        exit(1)
    if filepath.suffix not in FileExtensionSamples:
        print("File format not supported")
        exit(1)

def parse(filepath:Path)->list[dict]:
    exceptdetection(filepath)
    match filepath.suffix:
        case ".csv":
            return csv_parse(filepath)
        case ".json":
            return json_parse(filepath)
        case ".xlsx":
            return xlsx_parse(filepath)

def json_save(filepath:Path,data:list[dict]):
    dir_file = os.path.dirname(os.
path.abspath(filepath))
    fd, tmp = tempfile.mkstemp(dir=dir_file, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding = "utf-8") as f:
            json.dump(data,f,ensure_ascii=False, indent=2)
        os.replace(tmp,filepath)
    except Exception:
        os.remove(tmp)
        print("An error occurred while writing to the file.")
        raise

def csv_save(filepath:Path,data:list[dict]):
    dir_file = os.path.dirname(os.
path.abspath(filepath))
    fd, tmp = tempfile.mkstemp(dir=dir_file, suffix=".tmp")
    try:
        with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=_fieldnames(data))
            w.writeheader()
            w.writerows(data)
    except Exception:
        os.remove(tmp)
        print("An error occurred while writing to the file.")
        raise

def xlsx_save(path, data):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    names = _fieldnames(data)
    ws.append(names)
    for row in data:
        ws.append([row.get(k) for k in names])
    wb.save(path)

def _fieldnames(data):
    names = []
    for row in data:
        for k in row:
            if k not in names:
                names.append(k)
    return names

def save(filepath:Path,data:list[dict]):
    try:
        if not data:
            raise
    except Exception:
        print("The data to be written to the file is empty.")
        raise
    exceptdetection(filepath)
    match filepath.suffix:
        case ".csv":
            csv_save(filepath,data)
            return
        case ".json":
            json_save(filepath,data)
            return
        case ".xlsx":
            xlsx_save(filepath,data)
            return